"""Run 4 (28.1): download and prepare the UCI "Online Retail II" dataset.

Source, license (CC BY 4.0) and citation: docs/data-sources.md. Raw data is
cached OUTSIDE the repository and never committed.

What this script does, once, idempotently:
1. Download the zip (43.5 MB) to the cache dir, verify its SHA-256.
2. Extract `online_retail_II.xlsx` (two sheets: 2009-2010, 2010-2011).
3. Convert both sheets to ONE CSV (`online_retail_ii.csv`) with a
   deterministic per-row key `ROW_ID = <sheet>:<row index>` — the export has
   no stable line key, and the engine requires one. True duplicate rows in
   the source therefore become LOOK-ALIKE duplicates (flagged, not removed),
   which is the honest treatment per spec 5.3.
4. Derive `customers.csv` (customer_id, region = the customer's modal
   Country) as a reference table for segmentation context.

Usage:
    uv run python scripts/download_online_retail.py [--cache-dir DIR]
    (prints the prepared data directory path)
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import sys
import urllib.request
import zipfile
from pathlib import Path

URL = "https://archive.ics.uci.edu/static/public/502/online+retail+ii.zip"
# SHA-256 of the zip as first downloaded 2026-10-21; a mismatch means UCI
# changed the file and the mapping/case study must be re-verified.
EXPECTED_SHA256 = "572e36277c2390fbfde10664750731e0a86f55e33470d91919085f0408e67bfb"
XLSX_NAME = "online_retail_II.xlsx"
CSV_NAME = "online_retail_ii.csv"
CUSTOMERS_NAME = "customers.csv"
ITEMS_NAME = "items.csv"

HEADERS = [
    "Invoice",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "Price",
    "Customer ID",
    "Country",
]


def default_cache_dir() -> Path:
    base = os.environ.get("LOCALAPPDATA") or str(Path.home() / ".cache")
    return Path(base) / "margin-engine" / "online_retail_ii"


def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def download(cache: Path) -> Path:
    zip_path = cache / "online_retail_ii.zip"
    if not zip_path.exists():
        print(f"downloading {URL} ...", flush=True)
        tmp = zip_path.with_suffix(".part")
        req = urllib.request.Request(URL, headers={"User-Agent": "margin-engine/0.1"})
        with urllib.request.urlopen(req, timeout=120) as resp, open(tmp, "wb") as out:
            while True:
                chunk = resp.read(1 << 20)
                if not chunk:
                    break
                out.write(chunk)
        tmp.rename(zip_path)
    digest = sha256_of(zip_path)
    (cache / "online_retail_ii.zip.sha256").write_text(digest + "\n", encoding="utf-8")
    if EXPECTED_SHA256 != "REPLACED_AFTER_FIRST_DOWNLOAD" and digest != EXPECTED_SHA256:
        print(
            f"CHECKSUM MISMATCH: expected {EXPECTED_SHA256}, got {digest}. "
            "UCI changed the file; re-verify the mapping and case study before use.",
            file=sys.stderr,
        )
        sys.exit(1)
    print(f"zip sha256: {digest}")
    return zip_path


def convert(cache: Path, zip_path: Path) -> Path:
    csv_path = cache / CSV_NAME
    if csv_path.exists():
        print(f"already converted: {csv_path}")
        ensure_currency_column(csv_path)
        derive_reference(cache, csv_path)
        write_fx_rates(cache)
        return cache
    xlsx_path = cache / XLSX_NAME
    if not xlsx_path.exists():
        with zipfile.ZipFile(zip_path) as z:
            member = next(n for n in z.namelist() if n.endswith(".xlsx"))
            xlsx_path.write_bytes(z.read(member))

    from openpyxl import load_workbook

    print("converting xlsx -> csv (two sheets, ~1.07M rows; a few minutes)...", flush=True)
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    n = 0
    with open(csv_path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["ROW_ID", *HEADERS, "CURRENCY"])
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            rows = ws.iter_rows(values_only=True)
            header = [str(c) if c is not None else "" for c in next(rows)]
            if header != HEADERS:
                print(f"unexpected header in sheet {sheet}: {header}", file=sys.stderr)
                sys.exit(1)
            tag = sheet.replace(" ", "")
            for i, row in enumerate(rows, start=2):
                inv, stock, desc, qty, dt, price, cust, country = row
                cust_s = str(int(cust)) if isinstance(cust, float) else (str(cust) if cust else "")
                out = [
                    f"{tag}:{i}",
                    str(inv) if inv is not None else "",
                    str(stock) if stock is not None else "",
                    str(desc) if desc is not None else "",
                    str(qty) if qty is not None else "",
                    dt.strftime("%Y-%m-%d") if dt is not None else "",
                    f"{price:.4f}" if isinstance(price, (int, float)) else "",
                    cust_s,
                    str(country) if country is not None else "",
                    "GBP",  # the whole export is GBP (UK wholesaler)
                ]
                w.writerow(out)
                n += 1
    wb.close()
    print(f"wrote {n:,} lines to {csv_path}")

    derive_reference(cache, csv_path)
    write_fx_rates(cache)
    return cache


def ensure_currency_column(csv_path: Path) -> None:
    """One-time upgrade of an already-converted CSV (idempotent)."""
    import polars as pl

    head = csv_path.open(encoding="utf-8").readline()
    if "CURRENCY" in head:
        return
    df = pl.read_csv(csv_path, infer_schema=False).with_columns(pl.lit("GBP").alias("CURRENCY"))
    df.write_csv(csv_path)
    print("added constant CURRENCY=GBP column")


def write_fx_rates(cache: Path) -> None:
    """Monthly GBP->USD rates for the dataset's horizon, from the Frankfurter
    connector (offline via the committed fixture; docs/data-sources.md)."""
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
    from margin_engine.connectors.fx import fetch_daily_rates, monthly_rates, write_fx_rates_csv

    fixture = (
        Path(__file__).resolve().parents[1]
        / "tests"
        / "fixtures"
        / "frankfurter_gbp_usd_2009_2011.json"
    )
    daily, notes = fetch_daily_rates(
        "GBP", "USD", "2009-12-01", "2011-12-09", cache / "connector_cache", fixture=fixture
    )
    monthly = monthly_rates(daily, "GBP")
    write_fx_rates_csv(
        monthly,
        cache / "fx_rates.csv",
        columns={"currency": "CURRENCY", "period": "PERIOD", "rate": "RATE"},
    )
    for n in notes:
        print(n)
    print(f"wrote {monthly.height} monthly GBP->USD rates to {cache / 'fx_rates.csv'}")


def derive_reference(cache: Path, csv_path: Path) -> None:
    """customers.csv (modal country -> region) and items.csv (modal
    description per StockCode), derived from the converted CSV."""
    import polars as pl

    df = pl.read_csv(csv_path, infer_schema=False)
    customers = (
        df.filter(pl.col("Customer ID") != "")
        .group_by("Customer ID", "Country")
        .agg(pl.len().alias("n"))
        .sort(["Customer ID", "n", "Country"], descending=[False, True, False])
        .group_by("Customer ID", maintain_order=True)
        .agg(pl.col("Country").first())
        .rename({"Customer ID": "CUSTOMER_ID", "Country": "COUNTRY"})
        .sort("CUSTOMER_ID")
    )
    customers.write_csv(cache / CUSTOMERS_NAME)
    print(f"wrote {customers.height:,} customers to {cache / CUSTOMERS_NAME}")
    items = (
        df.filter(pl.col("Description") != "")
        .group_by("StockCode", "Description")
        .agg(pl.len().alias("n"))
        .sort(["StockCode", "n", "Description"], descending=[False, True, False])
        .group_by("StockCode", maintain_order=True)
        .agg(pl.col("Description").first())
        .rename({"StockCode": "ITEM_ID", "Description": "DESCRIPTION"})
        .sort("ITEM_ID")
    )
    items.write_csv(cache / ITEMS_NAME)
    print(f"wrote {items.height:,} items to {cache / ITEMS_NAME}")


def write_sample(cache: Path, out_dir: Path, n: int = 5000) -> None:
    """Committed demo sample (28.5): the first `n` lines of the converted CSV
    (deterministic; spans sales, returns and anonymous customers), with the
    matching reference rows. CC BY 4.0 with attribution — see the README
    written next to it."""
    import polars as pl

    out_dir.mkdir(parents=True, exist_ok=True)
    lines = pl.read_csv(cache / CSV_NAME, infer_schema=False).head(n)
    lines.write_csv(out_dir / CSV_NAME)
    custs = set(lines["Customer ID"].drop_nulls().to_list())
    items = set(lines["StockCode"].to_list())
    pl.read_csv(cache / CUSTOMERS_NAME, infer_schema=False).filter(
        pl.col("CUSTOMER_ID").is_in(sorted(custs))
    ).write_csv(out_dir / CUSTOMERS_NAME)
    pl.read_csv(cache / ITEMS_NAME, infer_schema=False).filter(
        pl.col("ITEM_ID").is_in(sorted(items))
    ).write_csv(out_dir / ITEMS_NAME)
    import shutil

    shutil.copy(cache / "fx_rates.csv", out_dir / "fx_rates.csv")
    readme = (
        "# Online Retail II - committed demo sample\n\n"
        f"The first {n:,} converted invoice lines of the UCI 'Online Retail II' "
        "dataset, with matching derived reference rows. Created by "
        "`scripts/download_online_retail.py --write-sample`.\n\n"
        "- Source: Chen, D. (2012). Online Retail II [Dataset]. UCI Machine "
        "Learning Repository. https://doi.org/10.24432/C5CG6D\n"
        "- License: CC BY 4.0 (https://creativecommons.org/licenses/by/4.0/). "
        "This sample is a format-converted excerpt; changes: two xlsx sheets "
        "converted to CSV, a deterministic ROW_ID and a constant CURRENCY=GBP "
        "column added, dates reduced to days.\n"
        "- fx_rates.csv holds monthly GBP->USD month-end reference rates "
        "derived from ECB daily rates via frankfurter.dev (see "
        "docs/data-sources.md).\n"
    )
    (out_dir / "README.md").write_text(readme, encoding="utf-8")
    print(f"sample written to {out_dir}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--cache-dir", type=Path, default=default_cache_dir())
    ap.add_argument("--write-sample", type=Path, default=None)
    args = ap.parse_args()
    cache: Path = args.cache_dir
    cache.mkdir(parents=True, exist_ok=True)
    zip_path = download(cache)
    out = convert(cache, zip_path)
    if args.write_sample is not None:
        write_sample(cache, args.write_sample)
    print(f"DATA_DIR={out}")


if __name__ == "__main__":
    main()
